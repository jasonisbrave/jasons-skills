#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""税务报告文件解析：.docx → report_text.md，.xlsx → sheets.json，并产出 files_manifest.json。

用法:
    python parse_tax.py <文件或目录> -o <outdir>

- 输入目录时递归收集 .docx/.doc/.xlsx/.xls（跳过 ~$ 临时文件）
- .doc/.xls 先经 legacy_convert 转换到 <outdir>/converted/，再按新格式解析
- .docx  → <outdir>/report_text.md（段落含 SRC 注释与样式名，表格渲染为 markdown 表格）
- .xlsx  → <outdir>/sheets.json（公式 + 缓存值双份，只收非空单元格）
"""

from __future__ import annotations

import argparse
import sys
import warnings
from datetime import datetime
from pathlib import Path

# Windows 控制台输出防乱码（Python 3.7+）
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

_HERE = Path(__file__).resolve().parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))

import legacy_convert  # noqa: E402
from _common import convert_cell_value, dump_json, setup_console  # noqa: E402

DOCX_EXTS = {".docx"}
XLSX_EXTS = {".xlsx"}
LEGACY_EXTS = {".doc", ".xls"}
SUPPORTED_EXTS = DOCX_EXTS | XLSX_EXTS | LEGACY_EXTS
SKIP_PREFIXES = ("~$",)  # Office 临时锁文件


# ---------------------------------------------------------------- 文件收集

def collect_files(target: Path) -> list[Path]:
    """输入文件或目录 → 支持的文件列表（已排序、跳过 ~$ 临时文件）。"""
    if target.is_file():
        if target.suffix.lower() not in SUPPORTED_EXTS:
            raise ValueError(
                f"不支持的文件类型: {target.suffix}（支持 {'/'.join(sorted(SUPPORTED_EXTS))}）"
            )
        return [target]
    if not target.is_dir():
        raise FileNotFoundError(f"路径不存在: {target}")
    files = [
        p for p in target.rglob("*")
        if p.is_file()
        and p.suffix.lower() in SUPPORTED_EXTS
        and not p.name.startswith(SKIP_PREFIXES)
    ]
    return sorted(files, key=lambda p: (str(p).lower(), str(p)))


# ---------------------------------------------------------------- .docx 解析

def docx_to_markdown(path: Path, display: str) -> tuple[str, int, int]:
    """提取全部段落（含样式名）与表格 → markdown 文本块。"""
    from docx import Document  # python-docx

    doc = Document(str(path))
    lines: list[str] = []
    for i, para in enumerate(doc.paragraphs, start=1):
        try:
            style = para.style.name if para.style is not None else ""
        except Exception:
            style = ""
        lines.append(f'<!-- SRC file="{display}" para={i} style="{style}" -->')
        lines.append(para.text)
    for j, table in enumerate(doc.tables, start=1):
        lines.append(f'<!-- SRC file="{display}" table={j} -->')
        for r, row in enumerate(table.rows):
            cells = [
                c.text.replace("|", "\\|").replace("\r\n", " ").replace("\n", " ")
                for c in row.cells
            ]
            lines.append("| " + " | ".join(cells) + " |")
            if r == 0 and cells:
                lines.append("|" + " --- |" * len(cells))
    return "\n".join(lines), len(doc.paragraphs), len(doc.tables)


# ---------------------------------------------------------------- .xlsx 解析

def parse_xlsx(path: Path) -> tuple[dict, int]:
    """加载两遍：data_only=False 取公式，data_only=True 取缓存值。

    返回 ({"sheets": {...}, "defined_names": [...], "external_links": [...]}, 单元格数)。
    """
    warnings.filterwarnings("ignore")  # 模板里的 Data Validation / Print area 告警与复核无关
    import openpyxl
    try:
        from openpyxl.worksheet.formula import ArrayFormula
    except Exception:  # 老版本 openpyxl 无此类型
        ArrayFormula = None  # type: ignore[assignment]

    wb_f = openpyxl.load_workbook(str(path), data_only=False)
    wb_v = openpyxl.load_workbook(str(path), data_only=True)
    try:
        sheets: dict = {}
        cell_count = 0
        for ws in wb_f.worksheets:
            ws_v = wb_v[ws.title]
            cells: dict = {}
            for row in ws.iter_rows():
                for c in row:
                    fval = c.value
                    if ArrayFormula is not None and isinstance(fval, ArrayFormula):
                        fval = fval.text  # 数组公式取其公式文本
                    vval = ws_v[c.coordinate].value
                    has_formula = isinstance(fval, str) and fval.startswith("=")
                    if not has_formula and (fval is None or fval == "") and (vval is None or vval == ""):
                        continue  # 只收非空单元格
                    entry: dict = {}
                    if has_formula:
                        entry["f"] = fval
                        entry["v"] = convert_cell_value(vval)  # 未计算过 → None（JSON null）
                    else:
                        entry["v"] = convert_cell_value(fval if fval not in (None, "") else vval)
                    cells[c.coordinate] = entry
            sheets[ws.title] = {"dims": ws.calculate_dimension(), "cells": cells}
            cell_count += len(cells)

        defined_names: list[str] = []
        try:
            defined_names = sorted(str(k) for k in wb_f.defined_names.keys())
        except Exception:
            pass

        external_links: list[str] = []
        for link in getattr(wb_f, "_external_links", None) or []:
            target = getattr(getattr(link, "file_link", None), "Target", None)
            if target:
                external_links.append(str(target))

        return {"sheets": sheets, "defined_names": defined_names, "external_links": external_links}, cell_count
    finally:
        wb_f.close()
        wb_v.close()


# ---------------------------------------------------------------- 主流程

def run(input_path: str | Path, outdir: str | Path) -> dict:
    """解析入口。返回 manifest dict（同时写出三个产物文件）。"""
    target = Path(input_path)
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    converted_dir = out / "converted"

    files = collect_files(target)
    if not files:
        raise ValueError(f"未找到 .docx/.doc/.xlsx/.xls 文件: {target}")

    # 重名文件（不同子目录）→ 用相对路径作显示名，保证唯一
    name_counts: dict[str, int] = {}
    for p in files:
        name_counts[p.name] = name_counts.get(p.name, 0) + 1
    used_keys: set[str] = set()

    def display_name(p: Path) -> str:
        if name_counts.get(p.name, 0) > 1:
            base = str(p)
            if target.is_dir():
                try:
                    base = str(p.relative_to(target))
                except ValueError:  # 转换产物不在 target 之下
                    pass
        else:
            base = p.name
        key, i = base, 2
        while key in used_keys:
            stem, dot, ext = base.rpartition(".")
            key = f"{stem}({i}).{ext}" if dot else f"{base}({i})"
            i += 1
        used_keys.add(key)
        return key

    sheets_data: dict = {}
    md_parts: list[str] = []
    manifest_files: list[dict] = []
    has_docx = has_xlsx = False

    for p in files:
        display = display_name(p)
        ext = p.suffix.lower()
        entry: dict = {"file": display, "path": str(p)}

        if ext in LEGACY_EXTS:
            entry["role"] = "legacy"
            print(f"  [转换] {p.name} ...")
            conv = legacy_convert.run(p, converted_dir)
            if conv.get("ok"):
                conv_path = Path(conv["output"])
                entry["converted_to"] = conv_path.name
                conv_display = display_name(conv_path)
                conv_entry: dict = {"file": conv_display, "path": str(conv_path), "role": "converted",
                                    "converted_from": display}
                try:
                    if conv_path.suffix.lower() in DOCX_EXTS:
                        block, n_para, n_table = docx_to_markdown(conv_path, conv_display)
                        md_parts.append(block)
                        conv_entry.update({"paragraphs": n_para, "tables": n_table})
                        has_docx = True
                    else:
                        data, n_cells = parse_xlsx(conv_path)
                        sheets_data[conv_display] = data
                        conv_entry.update({"sheets": list(data["sheets"].keys()), "cells": n_cells})
                        has_xlsx = True
                except Exception as e:
                    conv_entry["error"] = f"解析转换结果失败: {type(e).__name__}: {e}"
                manifest_files.append(conv_entry)
            else:
                entry["error"] = conv.get("error", "转换失败")
                print(f"  [WARN] {p.name} 转换失败: {entry['error']}")
            manifest_files.append(entry)
            continue

        if ext in DOCX_EXTS:
            entry["role"] = "report_docx"
            block, n_para, n_table = docx_to_markdown(p, display)
            md_parts.append(block)
            entry.update({"paragraphs": n_para, "tables": n_table})
            has_docx = True
        elif ext in XLSX_EXTS:
            entry["role"] = "form_xlsx"
            data, n_cells = parse_xlsx(p)
            sheets_data[display] = data
            entry.update({"sheets": list(data["sheets"].keys()), "cells": n_cells})
            has_xlsx = True
        manifest_files.append(entry)

    manifest = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "report_dir": str(target),
        "files": manifest_files,
    }
    dump_json(out / "files_manifest.json", manifest)
    if has_xlsx:
        dump_json(out / "sheets.json", sheets_data)
    if has_docx:
        (out / "report_text.md").write_text("\n".join(md_parts) + "\n", encoding="utf-8")

    # 控制台摘要
    print(f"\n[parse_tax] 共 {len(manifest_files)} 个解析对象（含转换产物）")
    for e in manifest_files:
        desc = e.get("role", "?")
        if "paragraphs" in e:
            desc += f" | 段落 {e['paragraphs']} 表格 {e['tables']}"
        if "sheets" in e:
            desc += f" | sheets {len(e['sheets'])} 单元格 {e.get('cells', 0)}"
        if "error" in e:
            desc += f" | 错误: {e['error']}"
        print(f"  - {e['file']}  [{desc}]")
    return manifest


def main() -> None:
    setup_console()
    parser = argparse.ArgumentParser(description="税务报告文件解析（docx/xlsx/doc/xls）")
    parser.add_argument("input", help="输入文件或目录")
    parser.add_argument("-o", "--outdir", required=True, help="输出目录")
    args = parser.parse_args()
    try:
        run(args.input, args.outdir)
    except SystemExit:
        raise
    except Exception as e:
        print(f"[ERROR] 解析失败: {type(e).__name__}: {e}", file=sys.stderr)
        print("提示: 请确认输入路径存在且包含 .docx/.doc/.xlsx/.xls 文件；"
              ".doc/.xls 需要本机可用的 Microsoft Office 或 LibreOffice。", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
