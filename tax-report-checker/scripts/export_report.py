#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""复核结果导出：results.json → Excel《复核结果.xlsx》+ Markdown《复核报告.md》。

用法:
    python export_report.py <results.json> -o <outdir>

Excel 7 个 sheet：摘要 / 报告勾稽 / 公式健康 / 数值校验 / 完整性 / 文本格式 / 检查项。
- 分类 sheet 只放未通过（passed=false）的条目；passed=true 的只进"检查项"（全量）。
Markdown 按问题（error）→ 存疑（warning）分组，末尾统计。severity 只用 error/warning/info。
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

# Windows 控制台输出防乱码（Python 3.7+）
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

_HERE = Path(__file__).resolve().parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))

from _common import count_severities, load_json, setup_console  # noqa: E402

CHECK_TYPES = ["报告勾稽", "公式健康", "数值校验", "完整性", "文本格式"]
HEADERS = ["规则", "严重程度", "描述", "应为值", "实际值", "差异", "位置", "原文摘录", "复核标记"]
FIELD_MAP: dict = {
    "规则": "rule_name", "严重程度": "severity", "描述": "description",
    "应为值": "expected", "实际值": "actual", "差异": "difference",
    "位置": "source_location", "原文摘录": "context",
}  # "复核标记" 留空给人工填写，不在映射中
SEV_ORDER = ["error", "warning", "info"]


# ---------------------------------------------------------------- Excel

def _style_header(ws, row: int, ncols: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    fill = PatternFill("solid", fgColor="D9E1F2")
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")


def _write_rows(ws, items: list[dict], title: str | None = None) -> None:
    from openpyxl.styles import Font
    start = 1
    if title:
        cell = ws.cell(row=1, column=1, value=title)
        cell.font = Font(bold=True)
        start = 2
    for j, h in enumerate(HEADERS, start=1):
        ws.cell(row=start, column=j, value=h)
    _style_header(ws, start, len(HEADERS))
    for i, item in enumerate(items, start=start + 1):
        for j, h in enumerate(HEADERS, start=1):
            key = FIELD_MAP.get(h)
            value = item.get(key, "") if key else ""  # "复核标记" 留空给人工
            ws.cell(row=i, column=j, value=value if value is not None else "")
    ws.freeze_panes = ws.cell(row=start + 1, column=1)


def _write_summary(ws, items: list[dict], meta: dict, source: str) -> None:
    from openpyxl.styles import Font
    ws.cell(row=1, column=1, value="税务报告复核结果").font = Font(bold=True, size=14)
    info = [
        ("输入文件", source),
        ("报告目录", meta.get("report_dir", "")),
        ("检查范围", meta.get("scope", "")),
        ("生成时间", meta.get("generated_at", datetime.now().isoformat(timespec="seconds"))),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)

    counts = count_severities(items)
    row = 8
    ws.cell(row=row, column=1, value="严重程度").font = Font(bold=True)
    ws.cell(row=row, column=2, value="数量").font = Font(bold=True)
    for i, (key, label) in enumerate([
        ("error", "error"), ("warning", "warning"), ("info", "info"), ("passed", "passed(相符)"),
    ], start=row + 1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=counts.get(key, 0))

    row += len(SEV_ORDER) + 3
    ws.cell(row=row, column=1, value="检查类型 × 严重程度").font = Font(bold=True)
    header = ["check_type", "error", "warning", "info", "passed"]
    for j, h in enumerate(header, start=1):
        ws.cell(row=row + 1, column=j, value=h)
    _style_header(ws, row + 1, len(header))
    i = row + 2
    for ctype in CHECK_TYPES:
        sub = [r for r in items if r.get("check_type") == ctype]
        ws.cell(row=i, column=1, value=ctype)
        for j, sev in enumerate(SEV_ORDER, start=2):
            ws.cell(row=i, column=j, value=sum(1 for r in sub if r.get("severity") == sev and not r.get("passed")))
        ws.cell(row=i, column=5, value=sum(1 for r in sub if r.get("passed")))
        i += 1
    ws.cell(row=i, column=1, value="合计（全量）")
    ws.cell(row=i, column=1).font = Font(bold=True)
    ws.cell(row=i, column=2, value=counts["error"])
    ws.cell(row=i, column=3, value=counts["warning"])
    ws.cell(row=i, column=4, value=counts["info"])
    ws.cell(row=i, column=5, value=counts["passed"])


def export_excel(items: list[dict], meta: dict, out_path: Path, source: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet()
    ws.title = "摘要"
    _write_summary(ws, items, meta, source)

    for ctype in CHECK_TYPES:
        sheet = wb.create_sheet(ctype)
        sub = [r for r in items if r.get("check_type") == ctype and not r.get("passed")]
        _write_rows(sheet, sub, title=f"{ctype}（未通过 {len(sub)} 项）")

    all_sheet = wb.create_sheet("检查项")
    _write_rows(all_sheet, items, title=f"检查项（全量 {len(items)} 项，含通过）")

    widths = {"A": 22, "B": 10, "C": 60, "D": 16, "E": 16, "F": 12, "G": 40, "H": 50, "I": 12}
    for sheet in wb.worksheets:
        if sheet.title == "摘要":
            continue
        for col, w in widths.items():
            sheet.column_dimensions[col].width = w

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))


# ---------------------------------------------------------------- Markdown

def export_markdown(items: list[dict], meta: dict, out_path: Path) -> None:
    counts = count_severities(items)
    lines: list[str] = ["# 税务报告复核报告", ""]
    if meta:
        lines.append(f"- 报告目录: {meta.get('report_dir', '-')}")
        lines.append(f"- 检查范围: {meta.get('scope', '-')}")
        lines.append(f"- 生成时间: {meta.get('generated_at', '-')}")
        lines.append("")

    def section(title: str, severity: str) -> None:
        sub = [r for r in items if r.get("severity") == severity and not r.get("passed")]
        lines.append(f"## {title}（{len(sub)} 项）")
        lines.append("")
        if not sub:
            lines.append("_无_")
            lines.append("")
            return
        for i, r in enumerate(sub, start=1):
            lines.append(f"### {i}. [{r.get('rule_name', '')}] {r.get('description', '')}")
            lines.append("")
            if r.get("source_location"):
                lines.append(f"- 位置: {r['source_location']}")
            pair = []
            if r.get("expected"):
                pair.append(f"应为 {r['expected']}")
            if r.get("actual"):
                pair.append(f"实际 {r['actual']}")
            if r.get("difference"):
                pair.append(f"差异 {r['difference']}")
            if pair:
                lines.append("- 数值: " + "，".join(pair))
            if r.get("context"):
                lines.append(f"- 原文摘录: {r['context']}")
            if r.get("evidence"):
                lines.append(f"- 依据: {r['evidence']}")
            lines.append("")

    section("一、问题（error）", "error")
    section("二、存疑（warning）", "warning")

    lines.append("## 三、统计")
    lines.append("")
    lines.append("| check_type | error | warning | info | passed |")
    lines.append("| --- | --- | --- | --- | --- |")
    for ctype in CHECK_TYPES:
        sub = [r for r in items if r.get("check_type") == ctype]
        lines.append("| {} | {} | {} | {} | {} |".format(
            ctype,
            sum(1 for r in sub if r.get("severity") == "error"),
            sum(1 for r in sub if r.get("severity") == "warning"),
            sum(1 for r in sub if r.get("severity") == "info"),
            sum(1 for r in sub if r.get("passed")),
        ))
    lines.append(f"| **合计** | {counts['error']} | {counts['warning']} | {counts['info']} | {counts['passed']} |")
    lines.append("")
    lines.append(f"> 严重程度仅使用 error / warning / info；passed 为核对相符项"
                 f"（info 中另含统计类条目）。")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---------------------------------------------------------------- 入口

def run(results_json: str | Path, outdir: str | Path) -> dict:
    src = Path(results_json)
    data = load_json(src)
    if isinstance(data, dict):
        meta = data.get("meta", {})
        items = data.get("results", [])
    elif isinstance(data, list):
        meta, items = {}, data
    else:
        raise ValueError("results.json 格式错误：应为数组或 {meta, results}")

    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    xlsx_path = out / "复核结果.xlsx"
    md_path = out / "复核报告.md"
    export_excel(items, meta, xlsx_path, str(src))
    export_markdown(items, meta, md_path)

    counts = count_severities(items)
    print(f"[export_report] Excel → {xlsx_path}")
    print(f"[export_report] Markdown → {md_path}")
    print(f"  error {counts['error']} / warning {counts['warning']} / info {counts['info']} / passed {counts['passed']}")
    return {"excel": str(xlsx_path), "markdown": str(md_path), "counts": counts}


def main() -> None:
    setup_console()
    parser = argparse.ArgumentParser(description="复核结果导出（Excel + Markdown）")
    parser.add_argument("results_json", help="run_check.py 或各 check_*.py 产出的 results.json")
    parser.add_argument("-o", "--outdir", required=True, help="输出目录")
    args = parser.parse_args()
    try:
        run(args.results_json, args.outdir)
    except SystemExit:
        raise
    except Exception as e:
        print(f"[ERROR] 导出失败: {type(e).__name__}: {e}", file=sys.stderr)
        print("提示: 请确认输入为 check/run_check 生成的 results.json（数组或 {meta, results}）。", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
